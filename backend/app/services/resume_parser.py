from __future__ import annotations

from io import BytesIO
from zipfile import BadZipFile, ZipFile

import fitz
from docx import Document


class ResumeParseError(Exception):
    """Safe parse error suitable for returning to a user."""


class UnsupportedResumeType(ResumeParseError):
    pass


def extract_resume_text(filename: str, content: bytes) -> str:
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix not in {"pdf", "docx"}:
        raise UnsupportedResumeType("Upload a PDF or DOCX file.")
    if not content:
        raise ResumeParseError("The uploaded file is empty.")

    try:
        if suffix == "pdf":
            if not content.startswith(b"%PDF-"):
                raise ResumeParseError("This file is not a readable PDF. Try exporting it again.")
            with fitz.open(stream=content, filetype="pdf") as document:
                text = "\n".join(page.get_text("text", sort=True) for page in document)
        else:
            if not content.startswith(b"PK"):
                raise ResumeParseError("This file is not a readable DOCX document.")
            with ZipFile(BytesIO(content)) as archive:
                expanded_size = sum(member.file_size for member in archive.infolist())
                if expanded_size > 15 * 1024 * 1024:
                    raise ResumeParseError("This DOCX expands beyond the safe parsing limit.")
            document = Document(BytesIO(content))
            text = "\n".join(
                paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()
            )
            for table in document.tables:
                for row in table.rows:
                    text += "\n" + " | ".join(cell.text for cell in row.cells)
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError(
            "The file could not be parsed. Please check it and try a text-based PDF or DOCX."
        ) from exc

    cleaned = "\n".join(line.strip() for line in text.splitlines() if line.strip()).strip()
    if len(cleaned) < 20:
        if suffix == "pdf":
            raise ResumeParseError(
                "No selectable text was found. This may be a scanned/image-only PDF; use a text-based PDF or paste the resume text."
            )
        raise ResumeParseError(
            "No usable text was found in this DOCX. Try another file or paste the resume text."
        )
    return cleaned
